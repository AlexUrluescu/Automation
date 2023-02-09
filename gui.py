from tkinter import *
from tkinter import filedialog
from tkinter.ttk import Combobox
import openpyxl
from email.message import EmailMessage
import ssl
import smtplib
from tkinter import messagebox
from tkinter import ttk
import time
import logging

logging.basicConfig(level=logging.DEBUG, format="%(asctime)s %(levelname)s %(message)s", datefmt="%Y-%m-%d %H:%M:%S")

# ----------------------------------------------------------
root = Tk()
root.title("Send students grades")
root.resizable(0,0)
root.geometry("600x650")

drop_down_sheet_list = StringVar()
drop_down_sheet_list.set("SheetName")

# --------------------------------------------------------

contor_students = 0
grades_var = IntVar()
absente_var = IntVar()
extra_text_var = IntVar()
restante_var = IntVar()

# --------------- Functions ------------------------------

def file_select_btn_clb():
    root.filename = filedialog.askopenfilename(initialdir="/", title="Select a file", filetypes = (("Excel files", "*.xlsx"),("All files", "*.*")))
    path_label.config(text=root.filename)
    path = root.filename
    combobox.configure(state=NORMAL)

    return path


def parameters_selected_btn_clb():
    path = (path_label.cget("text"))

    return path


def clear_widget(contor_students):
    info_students = f"S-au incarcat: {contor_students} studenti"
    info_students_data.delete(1.0, END)
    info_students_data.insert(1.0, info_students)
    info_students_data.configure(state=DISABLED)


def load_sheets_clb():
    path = parameters_selected_btn_clb()
    book = open_xlsx(path)
    combobox.config(values=book.sheetnames)
    


def open_xlsx(path):
    book = openpyxl.load_workbook(path, data_only=True)
    root.update()
    return book


def email_emisor():
    email = input_email.get()
    return email


def password_emisor():
    password = input_password.get()
    return password


def get_subject():
    subject = input_subject.get().capitalize()
    return subject


def get_sheet(book):
    list_data = []
    sheet = drop_down_sheet_list.get()
    book.active = book[sheet]
    book_active = book.active
    print(book_active.tables.items())
    paramas_table = book_active.tables.items()
    for data in paramas_table:
        list_data.append(data[1])

    list_data.append(book_active)
    
    return list_data


def get_table_data(sheet, range):
    global contor_students
    data_array = []
    array_students = []

    active_sheet = sheet
    print(active_sheet) 

    data_cells = active_sheet[range]
    for row in data_cells:
        data_array.append([cell_data.value for cell_data in row])
   

    for student in data_array:
        if(student[0] == None):
            break

        array_students.append(student)

    array_students.pop(0)

    print(array_students)

    
    contor_students = len(array_students)
    clear_widget(contor_students)
    print(contor_students)
    parameters_selected.configure(state=NORMAL)
    return array_students


def send_emails(data, email, password, subject, room, date):

    print(f"from send_emails_data -> {data}")

# progress bar placement in the root ----------------
    progress_bar.place(x=50, y=580) 
    progress_bar_label.place(x=230, y=550)

    range_progress = len(data) # variable for the dynamic increment

# ------------------------------------------------------

    for student in data:

        student_firstName = student[0]
        student_lastName = student[1]
        student_email = student[2]
        student_grade = student[3]
        student_abs = student[4]
        student_restanta = student[5]

        email_emisor = email
        email_password = password
        email_receptor = student_email

        if(grades_var.get() == 1 and absente_var.get() == 0):
            email_subject = 'Nota examen la ' + subject
            email_body = 'Buna ziua ' + student_firstName + " " + student_lastName + ', ai luat nota ' + str(student_grade) + "\n" + text_proba.get(1.0, END) + "."

        elif(absente_var.get() == 1 and grades_var.get() == 0):
            email_subject = 'Numarul de absente la ' + subject
            email_body = 'Buna ziua ' + student_firstName + " " + student_lastName + ', ai in total ' + str(student_abs) + " absente. \n" + text_proba.get(1.0, END) + "."

        elif(absente_var.get() == 1 and grades_var.get() == 1):
            email_subject = 'Notele si absentele la ' + subject
            email_body = 'Buna ziua ' + student_firstName + " " + student_lastName + ', ai nota ' + str(student_grade) + ' si ai in total ' + str(student_abs) + " absente. \n" + text_proba.get(1.0, END) + "."

        elif(restante_var.get() == 1):
            email_subject = "Informatii despre restanta la " + subject
            email_body = "Buna ziua " + student_firstName + " " + student_lastName + ", ai restanta in data de " + date +" in clasa " + room + ".\n" + text_proba.get(1.0, END) + "."

        em = EmailMessage()
        em['From'] = email_emisor
        em['To'] = email_receptor
        em['Subject'] = email_subject
        em.set_content(email_body)

        contexto = ssl.create_default_context()

        with smtplib.SMTP_SSL('smtp.gmail.com', 465, context=contexto) as smtp:
            smtp.login(email_emisor, email_password)
            smtp.sendmail(email_emisor, email_receptor, em.as_string())

    #progress bar dynamic increment ----------------------------------

        progress_bar['value'] += 100/range_progress
        progress_bar_label.config(text= str(progress_bar['value'])+"%")
        root.update_idletasks()
        time.sleep(1)

    # --------------------------------------------------------------------

        print('s-a terminat')


def get_data_btn_clb():
    path = parameters_selected_btn_clb()
    book = open_xlsx(path)
    range,sheet = get_sheet(book)
    email = email_emisor()
    print(email)
    password = password_emisor()
    print(password)
    get_table_data(sheet, range)

    # if(extra_text_var.get() == 1):
    #     text_proba.place(x=50, y=440)


def get_room():
    room = input_room.get()
    return room


def get_date():
    date = input_date.get()
    return date


def main():
    path = parameters_selected_btn_clb()
    book = open_xlsx(path)
    range,sheet = get_sheet(book)
    email = email_emisor()
    password = password_emisor()
    subject = get_subject()
    room = get_room()
    date = get_date()
    data = get_table_data(sheet, range)
    send_emails(data, email, password, subject, room, date)
    

def btn_info_clb():
    messagebox.showinfo("Information", "How can I generate the email email password: https://www.youtube.com/watch?v=DDVpKvJXRz8&list=PLdkIA_6OrXkLsQFuORCmRnyAEhO4Niiux&index=1")
    logging.debug("Butonul info a fost apasat")


def restante_callback():
    if(restante_var.get() == 1):
        input_room.config(state=NORMAL)
        input_date.config(state=NORMAL)
        absente_checkbutton.config(state=DISABLED)
        grades_checkbutton.config(state=DISABLED)
        grades_checkbutton.deselect()
        
        

    if(restante_var.get() == 0):
        input_room.config(state=DISABLED)
        input_date.config(state=DISABLED)
        absente_checkbutton.config(state=NORMAL)
        grades_checkbutton.config(state=NORMAL)

    logging.debug(f"{restante_var.get()}")
    logging.debug(f"grades: {grades_var.get()}")


def extraText_callback():
    if(extra_text_var.get() == 1):
        text_proba.place(x=50, y=440)

    if(extra_text_var.get() == 0):
        text_proba.place(x=1000, y=1000)
# -----------------------------------------------------------------

# ------------------- GUI -------------------------------------------

path_to_file = Label(root, text="Path")
path_to_file.place(x=50, y=20)

path_label = Label(root, text="")
path_label.place(x=85, y=22)
path_label.config(bg="black", width=40, fg="white")

sheet_label = Label(root, text="Sheet")
sheet_label.place(x=50, y=70)

email_label = Label(root, text="Email")
email_label.place(x=50, y=120)

input_email = Entry(root)
input_email.place(x=90, y=120)
input_email.config(width=35)

password_label = Label(root, text="Password")
password_label.place(x=50, y=160)

input_password = Entry(root)
input_password.place(x=130, y=160)

subject_label = Label(root, text="Subject")
subject_label.place(x=280, y=70)

input_subject = Entry(root)
input_subject.place(x=340, y=70)

details_author_label = Label(root, text="Made by Alexandre Urluescu, contact: alexurluescu23@gmail.com")
details_author_label.place(x=70, y=620)

# ---- entry for restante ---------

input_room = Entry(root, state=DISABLED)
input_room.place(x=450, y=210)

input_date = Entry(root, state=DISABLED)
input_date.place(x=450, y=240)

# ------- label restante  ------------

room_label = Label(root, text="Room")
room_label.place(x=400, y=210)

date_label = Label(root, text="Exam date")
date_label.place(x=370, y=240)

# ComboBox -----------------------------------------------------------------------------------
combobox = Combobox(root, textvariable= drop_down_sheet_list, postcommand=load_sheets_clb)
combobox.place(x=90, y=70)
combobox.configure(state=DISABLED)

# TextArea ----------------------------------------------------------------------------------

info_students_data = Text(root, width=50, height=5);
info_students_data.insert(1.0, "")
info_students_data.place(x=50, y=340)

text_proba = Text(root, width=50, height=5);
text_proba.insert(1.0, "")

# ------- Buttons -------------------------------------------------------------------------

file_select_btn = Button(root, text="Select file", command=file_select_btn_clb)
file_select_btn.place(x=400, y=20)

parameters_selected = Button(root, text="Send", width=15, command=main)
parameters_selected.place(x=250, y=280)
parameters_selected.configure(state=DISABLED)

get_data_btn = Button(root, text="Get students data", command=get_data_btn_clb)
get_data_btn.place(x=50, y=280)

btn_info = Button(root, text="Info", command=btn_info_clb)
btn_info.place(x=300, y=160)

# ----------------------------------------------------------------------------------

# --------- Checkbuttons ----------------------------------------------------------

grades_checkbutton = Checkbutton(root, text="Grades", variable=grades_var, onvalue=1, offvalue=0)
grades_checkbutton.place(x=50, y=230)
grades_checkbutton.select()

absente_checkbutton = Checkbutton(root, text="Absente", variable=absente_var, onvalue=1, offvalue=0)
absente_checkbutton.place(x=150, y=230)

extraText_checkbutton = Checkbutton(root, text="Extra text", variable=extra_text_var, onvalue=1, offvalue=0, command=extraText_callback)
extraText_checkbutton.place(x=250, y=230)

restante_checkbutton = Checkbutton(root, text="Restante", variable=restante_var, onvalue=1, offvalue=0, command=restante_callback)
restante_checkbutton.place(x=400, y=300)

# ----------------------------------------------------------------------------------------------

# ---------- Progress Bar --------------------------------------------------------------------

progress_bar = ttk.Progressbar(root, orient=HORIZONTAL, length=400, mode="determinate")

progress_bar_label = Label(root, text='')

# -------------------------------------------------------------------------------

root.mainloop()


if __name__ == "__main__":
    main()